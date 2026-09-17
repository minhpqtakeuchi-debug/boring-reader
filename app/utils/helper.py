from app.core.errors.exceptions.internal_server_error import InternalServerError

import re
import uuid
from app.utils.pdf_extractor.pdf_text_builder import PDFTextBuilder
from docx import Document as DocxDocument
import io
from typing import List, Optional, Dict, Any, Tuple
import csv


def camel_to_snake(name: str) -> str:
    """Convert camelCase or PascalCase to snake_case."""
    s1 = re.sub(r'(.)([A-Z][a-z]+)', r'\1_\2', name)
    s2 = re.sub(r'([a-z0-9])([A-Z])', r'\1_\2', s1)
    return s2.lower()


def is_valid_uuid(value: str) -> bool:
    try:
        uuid.UUID(value)
    except (ValueError, TypeError):
        return False
    return True

def _paragraph_block_json(
    text: str,
    *,
    doc_page: Optional[int] = None,
    bbox: Optional[Tuple[float, float, float, float]] = None
) -> Dict[str, Any]:
    b: Dict[str, Any] = {"type": "paragraph", "text": text}
    if doc_page is not None:
        b["doc_page"] = int(doc_page)
    if bbox is not None:
        b["bbox"] = [float(v) for v in bbox]
    return b

def _table_block_json_from_grid(
    grid: List[List[str]],
    *,
    doc_page: int = 0,
    table_id: int = 1,
    include_bbox: bool = False
) -> Dict[str, Any]:
    # Convert a simple 2D string grid into your table schema (no bboxes)
    rows_out: List[List[Dict[str, Any]]] = []
    for r, row in enumerate(grid):
        out_row: List[Dict[str, Any]] = []
        for c, val in enumerate(row):
            cell: Dict[str, Any] = {
                "r": r, "c": c,
                "rowspan": 1, "colspan": 1,
                "text": (val or "").strip(),
            }
            if include_bbox:
                # No geometry available — omit or set None. We'll omit for cleanliness.
                pass
            out_row.append(cell)
        rows_out.append(out_row)

    tbl: Dict[str, Any] = {
        "type": "table",
        "version": 1,
        "doc_page": int(doc_page),
        "table_id": int(table_id),
        "n_rows": len(grid),
        "n_cols": (len(grid[0]) if grid else 0),
        "rows": rows_out,
    }
    # No x_ticks / y_ticks / bbox when geometry is unknown
    return tbl

# -----------------------------------
# Main entry: bytes -> unified JSON
# -----------------------------------

def extract_binary_to_json(
    file_bytes: bytes,
    file_type: str
) -> Dict[str, Any]:
    """
    Return {"size": len(blocks), "blocks": [...]} with paragraph/table blocks unified
    across pdf | docx | txt | csv. Geometry (bbox) kept only when available.
    """
    ft = (file_type or "").lower().strip()
    blocks: List[Dict[str, Any]] = []

    if ft == "pdf":
        # High-fidelity path (tables + paragraphs + bbox)
        result = PDFTextBuilder.build_text_from_pdf_words(file_bytes)
        # Convert "version" top-level to "size" for consistency
        if "blocks" in result:
            return {"size": len(result["blocks"]), "blocks": result["blocks"]}
        else:
            return {"size": 0, "blocks": []}

    elif ft == "docx":
        from docx import Document as DocxDocument  # type: ignore
        doc = DocxDocument(io.BytesIO(file_bytes))
        page = 0
        table_id = 1

        # ---- 1) Paragraphs ----
        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue
            blocks.append({
                "type": "paragraph",
                "doc_page": page,
                "text": text,
            })

        # ---- 2) Tables ----
        for t in doc.tables:
            grid: List[List[str]] = []
            for row in t.rows:
                grid.append([cell.text.strip() for cell in row.cells])

            blocks.append({
                "type": "table",
                "doc_page": page,
                "table_id": table_id,
                "n_rows": len(grid),
                "n_cols": (len(grid[0]) if grid else 0),
                "rows": [
                    [
                        {
                            "r": r,
                            "c": c,
                            "rowspan": 1,
                            "colspan": 1,
                            "text": grid[r][c],
                        }
                        for c in range(len(row))
                    ]
                    for r, row in enumerate(grid)
                ],
            })
            table_id += 1

        return {"size": len(blocks), "blocks": blocks}

    elif ft in ("txt", "text"):
        # Split plain text file into paragraphs (blank-line separated)
        txt = file_bytes.decode("utf-8", errors="ignore")
        paras = [seg.strip() for seg in txt.replace("\r\n", "\n").split("\n\n") if seg.strip()]
        if not paras and txt.strip():
            paras = [txt.strip()]
        page = 0

        for para in paras:
            blocks.append({
                "type": "paragraph",
                "doc_page": page,
                "text": para,
            })

        return {"size": len(blocks), "blocks": blocks}

    elif ft == "csv":
        import csv
        s = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(s))
        grid = [list(row) for row in reader]
        page = 0
        table_id = 1

        blocks.append({
            "type": "table",
            "doc_page": page,
            "table_id": table_id,
            "n_rows": len(grid),
            "n_cols": (len(grid[0]) if grid else 0),
            "rows": [
                [
                    {
                        "r": r,
                        "c": c,
                        "rowspan": 1,
                        "colspan": 1,
                        "text": grid[r][c],
                    }
                    for c in range(len(row))
                ]
                for r, row in enumerate(grid)
            ],
        })

        return {"size": len(blocks), "blocks": blocks}

    else:
        raise InternalServerError(f"Unsupported file type: {file_type}")


def read_file_as_binary(file_path):
    """
    Reads a local file and returns its binary content.
    
    Args:
        file_path (str): Path to the file.
    
    Returns:
        bytes: Binary content of the file.
    """
    with open(file_path, "rb") as f:
        binary_data = f.read()
    return binary_data
